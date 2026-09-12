import os
import re
import datetime
import urllib.request
import numpy as np
import pandas as pd
import openpyxl

BASE_DIR = r"F:\stockModel"
DB_PATH = os.path.join(BASE_DIR, "jpx_daily_features_db.csv")
WEEKLY_DIR = os.path.join(BASE_DIR, "jpx_weekly_oi")
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# --- 1. 週次建玉Excelの自動探索 & ダウンロード ---
def fetch_latest_weekly_oi(target_date, file_type):
    os.makedirs(WEEKLY_DIR, exist_ok=True)
    year = target_date.strftime("%Y")

    # 直近の金曜日から過去2週間分を探索
    days_back = (target_date.weekday() - 4) % 7
    # 月曜日の公表前（17時前）ならさらに1週前を参照
    if target_date.weekday() == 0 and datetime.datetime.now().hour < 17:
        days_back += 7

    for offset in range(days_back, days_back + 15, 7):
        for day_shift in [0, 1, 2]: # 祝日考慮（金・木・水）
            ref_d = target_date - datetime.timedelta(days=offset + day_shift)
            ymd = ref_d.strftime("%Y%m%d")
            fname = f"{ymd}_{file_type}_oi_by_tp.xlsx"
            save_path = os.path.join(WEEKLY_DIR, fname)

            if os.path.exists(save_path) and os.path.getsize(save_path) > 1000:
                return save_path

            url = f"https://www.jpx.co.jp/automation/markets/derivatives/open-interest/files/{year}/{fname}"
            try:
                req = urllib.request.Request(url, headers=HEADERS)
                with urllib.request.urlopen(req) as resp:
                    with open(save_path, "wb") as f:
                        f.write(resp.read())
                print(f"  [+] 新規週次建玉ファイル取得成功 ({file_type}): {fname}")
                return save_path
            except Exception:
                continue

    return None

# --- 2. オプション確定壁パーサー ---
def parse_op_walls(filepath):
    if not filepath or not os.path.exists(filepath): return None, None
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        strike_oi = {'call': {}, 'put': {}}
        p_strike, c_strike = None, None

        for row in ws.iter_rows(values_only=True):
            if len(row) > 1 and row[0] == 1 and row[1] is not None:
                try: p_strike = float(row[1])
                except Exception: p_strike = None

            if len(row) > 11 and row[10] == 1 and row[11] is not None:
                try: c_strike = float(row[11])
                except Exception: c_strike = None

            if p_strike and (20000 <= p_strike <= 85000):
                s = float(row[4]) if len(row) > 4 and isinstance(row[4], (int, float)) else 0.0
                b = float(row[7]) if len(row) > 7 and isinstance(row[7], (int, float)) else 0.0
                if (s + b) > 0: strike_oi['put'][p_strike] = strike_oi['put'].get(p_strike, 0.0) + s + b

            if c_strike and (20000 <= c_strike <= 85000):
                s = float(row[14]) if len(row) > 14 and isinstance(row[14], (int, float)) else 0.0
                b = float(row[17]) if len(row) > 17 and isinstance(row[17], (int, float)) else 0.0
                if (s + b) > 0: strike_oi['call'][c_strike] = strike_oi['call'].get(c_strike, 0.0) + s + b

        c_wall = max(strike_oi['call'], key=strike_oi['call'].get) if strike_oi['call'] else None
        p_wall = max(strike_oi['put'], key=strike_oi['put'].get) if strike_oi['put'] else None
        return float(c_wall) if c_wall else None, float(p_wall) if p_wall else None
    except Exception as e:
        print(f"  [!] OP解析エラー: {e}")
        return None, None

# --- 3. CTA先物純建玉パーサー ---
def parse_fut_cta(filepath):
    if not filepath or not os.path.exists(filepath): return None
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        cta_brokers = ['ＡＢＮクリアリン証券', 'バークレイズ証券', 'ソシエテＧ証券']
        current_product = None
        cta_net = 0.0

        for row in ws.iter_rows(values_only=True):
            col0 = str(row[0]).strip() if row[0] is not None else ""
            if "＜日経225先物＞" in col0: current_product = 'large'; continue
            elif "＜日経225mini＞" in col0: current_product = 'mini'; continue
            elif "＜" in col0: current_product = 'other'; continue

            if current_product in ['large', 'mini']:
                scale = 1.0 if current_product == 'large' else 0.1
                seller = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                buyer = str(row[6]).strip() if len(row) > 6 and row[6] else ""
                s = float(row[4]) * scale if len(row) > 4 and isinstance(row[4], (int, float)) else 0.0
                b = float(row[7]) * scale if len(row) > 7 and isinstance(row[7], (int, float)) else 0.0

                if any(k in buyer for k in cta_brokers): cta_net += b
                if any(k in seller for k in cta_brokers): cta_net -= s

        return float(cta_net)
    except Exception as e:
        print(f"  [!] 先物解析エラー: {e}")
        return None

# --- 4. 日次同期メイン実行 ---
def update_daily():
    today = datetime.date.today()
    today_str = today.strftime("%Y-%m-%d")

    print("=" * 65)
    print(f"[*] 日次マクロデータベース同期実行: {today_str}")
    print("=" * 65)

    if not os.path.exists(DB_PATH):
        raise FileNotFoundError(f"[!] {DB_PATH} が見つかりません。")

    df_db = pd.read_csv(DB_PATH, index_col=0, parse_dates=True)
    last_valid = df_db.iloc[-1].to_dict()

    # 土日の場合は直近営業日（金曜）をターゲットにするかスキップ
    if today.weekday() >= 5:
        print(f"  [-] 本日は休場日（土日）のため、DB追記をスキップします。")
        return

    # 最新の週次確定ファイルを探索・取得
    p_op = fetch_latest_weekly_oi(today, "nk225op")
    p_fut = fetch_latest_weekly_oi(today, "indexfut")

    c_wall, p_wall = parse_op_walls(p_op)
    cta_net = parse_fut_cta(p_fut)

    # 取得できない場合は前営業日値を完全引き継ぎ（ffill）
    c_wall = c_wall or last_valid.get('call_oi_wall', 40000.0)
    p_wall = p_wall or last_valid.get('put_oi_wall', 38000.0)
    cta_net = cta_net if cta_net is not None else last_valid.get('cta_net_futures', 0.0)
    g_flip = (c_wall + p_wall) / 2.0

    print(f"  [+] 反映建玉水準: Call壁={c_wall:.0f} / Put壁={p_wall:.0f} / Flip={g_flip:.0f}")
    print(f"  [+] 反映CTA先物 : {cta_net:+.1f} 枚 (ラージ換算)")

    # DBのカラム（call_oi_wall, put_oi_wall, gamma_flip, cta_net_futures）と完全整合
    df_db.loc[pd.to_datetime(today_str)] = {
        'call_oi_wall': float(c_wall),
        'put_oi_wall': float(p_wall),
        'gamma_flip': float(g_flip),
        'cta_net_futures': float(cta_net)
    }

    df_db.sort_index(inplace=True)
    df_db.to_csv(DB_PATH)
    print(f"\n[★] {DB_PATH} の更新完了（総日数: {len(df_db)} 営業日）")
    print(df_db.tail(3))

if __name__ == "__main__":
    update_daily()