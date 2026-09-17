import os
import re
import json
import datetime
import urllib.request
import unicodedata
import openpyxl

# pipeline/配下からでもmodules/を解決できるようにプロジェクトルートをsys.pathへ追加
import sys
sys.path.insert(0, r"F:\stockModel")
from modules.constants import CTA_BROKER_KEYWORDS

BASE_DIR = r"F:\stockModel"
RAW_DIR = os.path.join(BASE_DIR, "jpx_raw_year")
OUT_JSON = os.path.join(BASE_DIR, "macro_flow_signal.json")

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# --- 1. 日次手口ファイルの自動探索 & ダウンロード ---
def fetch_daily_volume_file(target_date: datetime.date, is_jnet=False):
    os.makedirs(RAW_DIR, exist_ok=True)

    suffix = "_volume_by_participant_whole_day_J-NET.xlsx" if is_jnet else "_volume_by_participant_whole_day.xlsx"

    # 当日から過去7日分を遡って探索（土日・祝日・公表前の対応）
    for back in range(7):
        ref_d = target_date - datetime.timedelta(days=back)
        # 土日はスキップ
        if ref_d.weekday() >= 5:
            continue

        ym = ref_d.strftime("%Y%m")
        ymd = ref_d.strftime("%Y%m%d")
        fname = f"{ymd}{suffix}"
        save_path = os.path.join(RAW_DIR, fname)

        # 既に有効なファイル（5KB以上）があればそれを利用
        if os.path.exists(save_path) and os.path.getsize(save_path) > 5000:
            return save_path, ymd

        url = f"https://www.jpx.co.jp/automation/markets/derivatives/participant-volume/files/daily/{ym}/{fname}"
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=10) as resp:
                content = resp.read()
                # 404などのHTMLエラーページ除外
                if b"<!DOCTYPE html" in content[:150] or b"<html" in content[:150]:
                    continue
                if len(content) < 5000:
                    continue

                with open(save_path, "wb") as f:
                    f.write(content)
            print(f"  [+] 日次手口ファイル取得成功: {fname}")
            return save_path, ymd
        except Exception:
            continue

    return None, None

# --- 2. エクセルから日経225先物の出来高を集計 ---
def extract_nk225_volumes(filepath, target_cta):
    if not filepath or not os.path.exists(filepath):
        return 0.0, 0.0
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        total_vol = 0.0
        cta_vol = 0.0

        for r in range(9, ws.max_row + 1):
            p_class = str(ws.cell(r, 1).value or "").strip()
            # NK225ラージ先物（NK225F）を主軸に集計
            if p_class == "NK225F":
                broker = unicodedata.normalize('NFKC', str(ws.cell(r, 6).value or "")).strip()
                v = ws.cell(r, 8).value
                vol = float(v) if isinstance(v, (int, float)) else 0.0
                total_vol += vol

                if any(k in broker for k in target_cta):
                    cta_vol += vol

        return total_vol, cta_vol
    except Exception as e:
        print(f"  [!] ファイル読み込みエラー ({os.path.basename(filepath)}): {e}")
        return 0.0, 0.0

# --- 3. メイン解析実行 ---
def analyze_daily_flow():
    today = datetime.date.today()
    print("=" * 65)
    print(f"[*] 日次大口手口（CTA/J-NET）フロー解析開始: {today.strftime('%Y-%m-%d')}")
    print("=" * 65)

    target_cta = CTA_BROKER_KEYWORDS

    # 立会（Auction）とJ-NETの両方を自動探索・ダウンロード
    f_auc, ymd_auc = fetch_daily_volume_file(today, is_jnet=False)
    f_jnet, ymd_jnet = fetch_daily_volume_file(today, is_jnet=True)

    if not f_auc:
        print("[!] 解析対象の立会手口ファイルが取得できませんでした。")
        return

    auc_tot, auc_cta = extract_nk225_volumes(f_auc, target_cta)
    jnet_tot, jnet_cta = extract_nk225_volumes(f_jnet, target_cta)

    combined_tot = auc_tot + jnet_tot
    combined_cta = auc_cta + jnet_cta

    cta_share = (combined_cta / combined_tot) if combined_tot > 0 else 0.0
    jnet_ratio = (jnet_tot / combined_tot) if combined_tot > 0 else 0.0

    # シグナル判定
    level = "NORMAL"
    badge_text = f"CTA: {(cta_share * 100):.1f}% | J-NET: {(jnet_ratio * 100):.1f}%"
    signal_desc = "機関・大口フロー平常"

    if cta_share >= 0.40:
        level = "CTA_SURGE"
        signal_desc = "⚡ CTA手口急増（トレンド加速／逆張り警戒）"
    elif jnet_ratio >= 0.25:
        level = "JNET_HIGH"
        signal_desc = "📦 機関大口クロス活発（ポジ移動／ロール注意）"
    elif cta_share < 0.20 and combined_tot > 0:
        level = "QUIET"
        signal_desc = "💤 大口薄商い（ダマシ警戒／保ち合い気味）"

    result = {
        "date": ymd_auc,
        "cta_share": round(cta_share, 3),
        "jnet_ratio": round(jnet_ratio, 3),
        "level": level,
        "badge_text": badge_text,
        "signal_desc": signal_desc
    }

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"  [+] 基準取引日  : {ymd_auc}")
    print(f"  [+] CTA出来高比率: {(cta_share * 100):.1f}% (合算 {combined_cta:.0f} / {combined_tot:.0f} 枚)")
    print(f"  [+] J-NET比率   : {(jnet_ratio * 100):.1f}%")
    print(f"  [★] 判定シグナル: {signal_desc}")
    print(f"[+] 保存完了 -> {OUT_JSON}")

if __name__ == "__main__":
    analyze_daily_flow()